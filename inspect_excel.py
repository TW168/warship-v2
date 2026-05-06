#!/usr/bin/env python3
"""Quick script to inspect the Excel file and see what columns it contains."""

import openpyxl

# Load the Excel file
wb = openpyxl.load_workbook('/home/tony/cfp/warship-v2/raw/2026-05-01-REPORT.xlsx', data_only=True)
ws = wb.active

print("Excel file inspection:")
print(f"Sheet name: {ws.title}")
print(f"Max row: {ws.max_row}")
print(f"Max column: {ws.max_column}")

# Get headers (first row)
headers = []
for col in range(1, ws.max_column + 1):
    header = ws.cell(row=1, column=col).value
    headers.append(str(header) if header is not None else f"col_{col}")

print(f"\nHeaders ({len(headers)} columns):")
for i, header in enumerate(headers, 1):
    print(f"  {i}. {header}")

# Show first few data rows
print(f"\nFirst 5 data rows:")
for row_num in range(2, min(7, ws.max_row + 1)):
    row_data = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=row_num, column=col).value
        row_data.append(str(cell_value) if cell_value is not None else "")
    print(f"  Row {row_num}: {row_data}")

wb.close()